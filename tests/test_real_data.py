from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]

from ladle_preallocation.real_data.pipeline import filter_plan_records, plan_key
from ladle_preallocation.real_data.reader import location_mapping_statistics, read_location
from ladle_preallocation.real_data.audit import LocationAwareAuditError, load_location_aware_audit
from ladle_preallocation.real_data.scenario import build_scenario_from_real_snapshots
from ladle_preallocation.real_data.xlsx_stream import iter_crane_rows, read_latest_online_cranes_stream

WORKBOOK = '''<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="CRANE_STATUS_HIS - 副本" sheetId="1" r:id="rId1"/></sheets></workbook>'''
RELS = '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>'''


def _cell(column: str, row: int, value: str | int | float | None) -> str:
    return "" if value is None else f'<c r="{column}{row}" t="inlineStr"><is><t>{value}</t></is></c>'


def _crane_row(row: int, crane: str, updated: str, online: int = 1, position: int = 10000) -> str:
    values = [row, crane, position, 0, 0, "A", online, updated, 2, 300, 30, 45, 48, 10, 0, 48, 0, 1, 1, None]
    cells = []
    for index, value in enumerate(values):
        number, label = index + 1, ""
        while number:
            number, remainder = divmod(number - 1, 26)
            label = chr(65 + remainder) + label
        cells.append(_cell(label, row, value))
    return f'<row r="{row}">{"".join(cells)}</row>'


def _fixture(path: Path) -> None:
    rows = [
        '<row r="1"><c r="A1" t="inlineStr"><is><t>header</t></is></c></row>',
        _crane_row(2, "3170", "2026-03-06-08.00.00.000000", 1, 10000),
        _crane_row(3, "3170", "2026-03-06-09.00.00.000000", 1, 12000),
        _crane_row(4, "3180", "2026-03-06-08.30.00.000000", 0, 20000),
        _crane_row(5, "3190", "2025-03-24-08.30.00.000000", 1, 30000),
    ]
    sheet = '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + "".join(rows) + '</sheetData></worksheet>'
    with ZipFile(path, "w") as archive:
        archive.writestr("xl/workbook.xml", WORKBOOK)
        archive.writestr("xl/_rels/workbook.xml.rels", RELS)
        archive.writestr("xl/worksheets/sheet1.xml", sheet)


class RealDataTests(unittest.TestCase):
    def test_location_aware_audit_is_required_for_replay(self) -> None:
        audit = load_location_aware_audit(ROOT / "outputs/real_data_location_aware/decision_tree_audit.json")
        self.assertEqual(audit["location"]["baseline_contract"], "location-aware-decision-tree-v1")
        target = next(item for item in audit["scheduling_inputs"]["ladles"] if item["ladle_id"] == "ST38")
        self.assertEqual(target["position_m"], 17230.0)
        assignment = next(item for item in audit["assignments"] if item["heat_id"] == "JU6310E7-300770")
        self.assertEqual((assignment["ladle_id"], assignment["crane_id"]), ("ST36", "4170"))
        with self.assertRaises(LocationAwareAuditError):
            load_location_aware_audit(ROOT / "outputs/real_data_validation/decision_tree_audit.json")

    def test_stream_reader_keeps_latest_online_snapshot(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "CRANE.xlsx"
            _fixture(path)
            self.assertEqual(len(list(iter_crane_rows(path))), 4)
            snapshots, stats = read_latest_online_cranes_stream(path, 1772780000, 1772790000)
            self.assertEqual([row["crane_id"] for row in snapshots], ["3170"])
            self.assertEqual(snapshots[0]["position_x"], "12000")
            self.assertEqual(stats.source_rows_scanned, 4)

    def test_plan_filter_uses_dominant_month_and_explains_exclusion(self) -> None:
        rows = [
            {"heat_id": "H1", "plan_sequence": 1, "tap_finish_at": "2026-03-06-08.00.00", "location_mapping_status": "已匹配"},
            {"heat_id": "H2", "plan_sequence": 2, "tap_finish_at": "2026-03-07-08.00.00", "location_mapping_status": "已匹配"},
            {"heat_id": "H3", "plan_sequence": 3, "tap_finish_at": "2025-03-24-08.00.00", "location_mapping_status": "已匹配"},
        ]
        included, excluded, month = filter_plan_records(rows)
        self.assertEqual(month, "2026-03")
        self.assertEqual(len(included), 2)
        self.assertIn("不在评测月份", excluded["H3::3"])
        self.assertEqual(plan_key(rows[0]), "H1::1")

    def test_location_reader_preserves_real_coordinates_and_metadata(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "loc_location.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append(["LOC_LOCATION", "SPAN_NAME", "POS_X", "LOC_TYPE", "LOC_STATUS", "GWDMC"])
            sheet.append(["4QF5", 3, 36075, None, 0, "4#倾翻台"])
            sheet.append(["2QF5", 3, 17230, None, 0, "2#倾翻台"])
            workbook.save(path)
            locations = read_location(path)
            self.assertEqual(locations["4QF5"]["pos_x"], 36075.0)
            self.assertEqual(locations["2QF5"]["description"], "2#倾翻台")
            self.assertTrue(locations["4QF5"]["position_valid"])

    def test_location_reader_rejects_duplicate_codes(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "loc_location.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append(["LOC_LOCATION", "SPAN_NAME", "POS_X", "LOC_TYPE", "LOC_STATUS", "GWDMC"])
            sheet.append(["4QF5", 3, 36075, None, 0, "one"])
            sheet.append(["4QF5", 3, 36076, None, 0, "two"])
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "重复 LOC_LOCATION"):
                read_location(path)

    def test_location_reader_marks_out_of_range_coordinates_invalid(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "loc_location.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append(["LOC_LOCATION", "SPAN_NAME", "POS_X", "LOC_TYPE", "LOC_STATUS", "GWDMC"])
            sheet.append(["BAD", 3, 48001, None, 0, "bad"])
            workbook.save(path)
            locations = read_location(path)
            self.assertFalse(locations["BAD"]["position_valid"])
            self.assertEqual(locations["BAD"]["position_invalid_reason"], "outside_coordinate_range")

    def test_mapped_position_is_used_and_missing_codes_are_audited(self) -> None:
        plan = [{
            "heat_id": "H1", "plan_sequence": 1, "allocation_flag": 1,
            "required_grade": "0", "decarb_ladle_id": "L1",
            "current_ladle_grade": "22A", "ladle_position": "4QF5",
            "empty_ladle_weight": 140000, "tap_finish_at": "2026-03-06-08.00.00",
            "ladle_arrival_at": "2026-03-06-07.50.00", "ladle_pour_finish_at": None,
        }, {
            "heat_id": "H2", "plan_sequence": 2, "allocation_flag": 1,
            "required_grade": "0", "decarb_ladle_id": "L2",
            "current_ladle_grade": "22A", "ladle_position": "UNKNOWN",
            "empty_ladle_weight": 140000, "tap_finish_at": "2026-03-06-08.10.00",
            "ladle_arrival_at": "2026-03-06-08.00.00", "ladle_pour_finish_at": None,
        }]
        locations = {"4QF5": {"location_code": "4QF5", "pos_x": 36075.0, "position_valid": True}}
        heats, ladles, _ = build_scenario_from_real_snapshots(
            plan,
            [{"crane_id": "C1", "position_x": 30000, "online_status": 1, "updated_at": "2026-03-06-08.00.00", "speed_mps": 2, "safe_distance_m": 10, "limit_0_m": 0, "limit_1_m": 48, "weight_tonnes": 0}],
            location_map=locations,
        )
        by_ladle = {row["ladle_id"]: row for row in ladles}
        self.assertEqual(by_ladle["L1"]["position_m"], 36075.0)
        self.assertEqual(by_ladle["L1"]["location_mapping_status"], "mapped")
        self.assertIsNone(by_ladle["L2"]["position_m"])
        self.assertEqual(by_ladle["L2"]["location_mapping_status"], "missing_location_code")
        stats = location_mapping_statistics(plan, locations)
        self.assertEqual(stats["mapped"], 1)
        self.assertEqual(stats["missing_location_code"], 1)


if __name__ == "__main__":
    unittest.main()
