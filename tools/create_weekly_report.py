from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/admin/Desktop/钢包配包")
TEMPLATE = Path("/Users/admin/Desktop/数据文件/2026-9-22 钢包重排可视化展示.docx")
OUTPUT = ROOT / "outputs/2026-9-23 钢包扰动重调度双线路评价.docx"
FONT_CN = "宋体"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="B7C9DC", size="8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, *, bold=False, color="000000", size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = FONT_CN
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def add_table(doc, headers, rows, widths, *, font_size=9, header_fill="24557D"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], header_fill)
        set_cell_border(table.rows[0].cells[i], "B7C9DC", "8")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_border(cells[i], "B7C9DC", "8")
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    return table


def add_summary_table(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].cells[0].width = Inches(6.5)
    cell = table.rows[0].cells[0]
    set_cell_text(cell, text, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_shading(cell, "E9F2F9")
    set_cell_border(cell, "A8C7E5", "8")
    return table


def add_body(doc, text, *, bold=False, color="000000", size=10.5, after=5, indent=True):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.5
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT_CN
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.font.name = FONT_CN
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r.font.size = Pt(16 if level == 1 else 13.5)
    r.bold = True
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def clear_document_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    for name, size in (("Heading 1", 16), ("Heading 2", 13.5)):
        style = doc.styles[name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))
    clear_document_body(doc)
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.83)
    section.right_margin = Inches(0.83)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    title = doc.add_paragraph(style="Normal")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("钢包扰动重调度双线路评价")
    r.font.name = FONT_CN
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r.font.size = Pt(28)
    r.bold = True
    doc.add_paragraph(style="Normal")

    add_heading(doc, "1. 本周结论")
    add_body(doc, "本周用现有的生产计划、行车和位置数据，模拟了 20 类生产扰动。每个场景都从同一份扰动前预配结果出发：一条线路由决策树重新配包，另一条线路由 Codex 独立重新配包，最后再用同一套规则检查结果是否完整、安全、可以执行。")
    add_body(doc, "实验结果是：决策树线路没有一个场景完整通过检查，Codex 线路 20 个场景全部通过。这里的 0/20 并不是决策树完全没有给出结果，而是每个场景都至少存在一个问题，例如有炉次没有完成分配，或者整体方案没有通过统一约束检查。")
    add_summary_table(doc, "已完成：20 类扰动双线路实验。决策树通过 0/20，Codex 通过 20/20。当前结果说明 Codex 能在本轮模拟场景中补上决策树没有解决的问题，但还需要现场数据进一步验证。")

    add_heading(doc, "2. 双线路评价方案")
    add_body(doc, "为了便于查看 Excel，本报告继续使用 AP、AQ 和 AR 三个简称：AP 是事故发生前的预配结果，AQ 是事故发生后由决策树给出的重排结果，AR 是事故发生后由 Codex 给出的重排结果。")
    add_table(doc, ["项目", "AQ 决策树线路", "AR Codex 线路"], [
        ["共同起点", "AP 扰动前预配结果", "AP 扰动前预配结果"],
        ["输入信息", "同一个事故状态和可用资源", "同一个事故状态和可用资源"],
        ["重新配包方式", "按固定规则局部重排", "根据约束独立生成重排结果"],
        ["是否读取另一线路", "否", "否"],
        ["结果检查", "使用同一套检查规则", "使用同一套检查规则"],
    ], [1.35, 2.55, 2.6], font_size=9)
    add_body(doc, "比较时先看方案能不能完整、安全地执行，包括是否每个受影响炉次都有钢包和行车、是否使用了可用资源、是否违反限制、是否改动了影响范围外的炉次。只有两条线路都能执行时，才继续比较按时率、延迟时间和改配数量。")

    page_break(doc)
    add_heading(doc, "3. 批量扰动实验结果")
    add_body(doc, "本轮共设置 20 类扰动。生产计划、炉次、钢包、行车和位置来自现有数据；事故状态由实验模拟。除一个 31 炉的行车离线主场景外，其余 19 个场景各选取 2 个相邻炉次，检查局部事故发生后两条线路能否重新完成配包。")
    add_table(doc, ["结果指标", "AQ 决策树", "AR Codex", "解读"], [
        ["完整通过检查的场景", "0/20", "20/20", "Codex 完成了全部场景"],
        ["31 炉主场景完成数", "30/31", "31/31", "Codex 补上了决策树未完成的 1 炉"],
        ["最终推荐场景数", "0", "20", "先判断能否执行，再决定推荐"],
        ["炉次结果记录", "69 条", "69 条", "两条线路对同一批炉次进行比较"],
        ["涉及的不同炉次", "33 炉", "33 炉", "主场景和局部场景合计"],
    ], [1.45, 1.15, 1.15, 2.75], font_size=8.7)

    add_heading(doc, "3.1 20 类扰动覆盖范围", level=2)
    add_table(doc, ["扰动类别", "覆盖场景", "数量"], [
        ["设备", "行车离线、速度下降、限载、作业范围受限", "4"],
        ["钢包", "钢包不可用、损坏、内衬报警、超龄、等级冲突", "5"],
        ["工艺设施", "精炼设施停机、氩站不可用、精炼路径变更", "3"],
        ["计划时序", "生产计划偏差、转炉延迟、连铸机延迟、紧急炉次插入、优先级提升", "5"],
        ["物流与安全", "运输通道封锁", "1"],
        ["数据控制", "行车遥测陈旧", "1"],
        ["复合扰动", "多资源复合扰动", "1"],
    ], [1.25, 4.25, 0.95], font_size=8.7)
    add_body(doc, "与前期只模拟行车离线相比，本轮已经覆盖行车、钢包、工艺设施、生产计划、运输通道、数据异常和复合事故。需要注意，这 20 类是为了测试系统而设计的模拟场景，并不是工厂历史事故的完整统计。")

    add_heading(doc, "3.2 31 炉行车离线主场景", level=2)
    add_body(doc, "主场景模拟 11:00 一台行车离线，事故影响窗口内共有 31 炉。决策树完成了其中 30 炉，还有 1 炉没有安排完成，因此整套结果不能直接执行。Codex 完成了 31 炉的重新安排，并通过全部检查。与事故前计划相比，Codex 调整了 3 炉的钢包或行车，另有 1 炉只调整了精炼路线。")
    add_table(doc, ["主场景观察", "AQ", "AR Codex"], [
        ["受影响窗口", "31 炉", "31 炉"],
        ["完成分配", "30/31", "31/31"],
        ["统一规则检查", "未全部通过", "全部通过"],
        ["结果说明", "还剩 1 炉未完成，整套方案不能直接执行", "31 炉全部完成，只调整必要的钢包、行车或路线"],
    ], [2.2, 2.0, 2.25], font_size=9)

    page_break(doc)
    add_heading(doc, "4. 工厂方案对照发现")
    add_body(doc, "PLAN 数据中有工厂原来的钢包选择，可以用来核对预配结果。但是当前还没有工厂的行车任务、事故发生后的人工重排结果、现场实际执行时间和验收结果，因此暂时不能把本轮算法结果与完整的人工处置方案直接比较。")
    add_table(doc, ["对照项目", "本轮结果", "应如何解释"], [
        ["有效工厂钢包记录", "436 条（PLAN 469 行中）", "可用于预配钢包字段对照"],
        ["可比场景炉次观察", "63 条，27 个唯一炉次", "仅限有工厂钢包记录的场景观察"],
        ["AP / AQ / AR 与工厂钢包一致率", "均为 0%", "说明算法使用的信息或选择规则与工厂仍有明显差距"],
        ["工厂天车与事故后人工方案", "暂无", "无法比较人工调度、资源可执行性和现场接受结果"],
    ], [2.15, 1.65, 2.65], font_size=8.8)
    add_body(doc, "一致率为 0% 不代表算法结果一定不能执行，也不能说明 Codex 已经优于工厂人员。它说明当前算法还没有掌握工厂选择钢包时使用的全部信息，例如现场实时状态和未写入数据表的经验规则。下一步要先补齐这些信息，再分析为什么选择不同。")

    add_heading(doc, "5. 证据边界与下一步")
    add_table(doc, ["当前可以下的结论", "当前不能下的结论"], [
        ["在本轮 20 类模拟扰动中，Codex 全部通过，决策树均未完整通过。", "不能把 20 类模拟场景说成工厂历史事故的完整集合。"],
        ["使用相同数据和检查规则时，Codex 补上了决策树没有完成的部分。", "不能据此宣称已经替代人工调度或完成现场生产验证。"],
        ["本轮结果可以重复生成，便于下一阶段继续核对。", "不能把离线运行时间当作真实接口的响应时间。"],
    ], [3.35, 3.1], font_size=8.8)
    add_table(doc, ["下一步补充数据", "用途"], [
        ["事故后的人工重排记录", "与现场人员的实际处理结果进行比较"],
        ["工厂行车任务和实际执行时间", "核对行车安排、等待时间和实际延迟"],
        ["钢包状态、包龄和等级数据", "检查钢包是否真的可用，并解释选择差异"],
        ["工厂内部配包和事故处置规则", "补上数据表中没有写明的现场经验"],
        ["实时 Codex 接口耗时", "确认能否在 90 秒内返回，并统计平均值和高峰耗时"],
    ], [3.0, 3.45], font_size=8.8)
    add_body(doc, "下一阶段先接入真实的人工重排和行车执行记录，再用相同方法重新回放。实时 Codex 接口也继续保留，重点检查它能否在 90 秒内稳定返回结果。")

    # Make the document metadata explicit for later audit.
    doc.core_properties.title = "钢包扰动重调度双线路评价"
    doc.core_properties.subject = "2026 年 9 月周报"
    doc.core_properties.author = ""
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
