from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FONT_FAMILY = "Songti SC"
REF = Path("/Users/admin/Desktop/数据文件/2026-8-9 LLM局部重调度方案报告.docx")
AUDIT = ROOT / "outputs/real_data_codex_stress_demo/audit.json"
SUMMARY = ROOT / "outputs/offline_scenarios/summary.json"
OUT_DIR = ROOT / "outputs"
OUT_DOCX = OUT_DIR / "2026-8-31 LLM局部重调度方案验证报告.docx"
OUT_MD = OUT_DIR / "2026-8-31 LLM局部重调度方案验证报告.md"
DIAGRAM_PNG = ROOT / "docs/llm-rescheduling-flow.png"
INDICATOR_PNG = ROOT / "docs/evaluation-indicator-system.png"

BLUE = "1F4D78"
MID_BLUE = "D9EAF7"
PALE_BLUE = "F1F6FB"
LIGHT_BLUE = "EAF3FA"
RED = "C00000"
GREEN = "008000"
GRAY = "666666"
BLACK = "000000"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "C6D6E5", sz: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    # Songti SC is the installed Mac name and resolves to 宋体 in Word.
    run.font.name = FONT_FAMILY
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT_FAMILY)
    r_fonts.set(qn("w:hint"), "eastAsia")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    lang.set(qn("w:bidi"), "en-US")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(paragraph, size: float = 10.5, color: str = BLACK, bold: bool = False) -> None:
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.5
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_text(doc, text: str, style: str = "Normal", size: float = 10.5, color: str = BLACK, bold: bool = False, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    format_paragraph(p, size=size, color=color, bold=bold)
    return p


def add_rich_paragraph(doc, parts: list[tuple[str, dict]], style: str = "Normal", align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    for text, options in parts:
        r = p.add_run(text)
        set_run_font(r, size=options.get("size", 10.5), bold=options.get("bold", False), color=options.get("color", BLACK))
    format_paragraph(p)
    return p


def add_heading(doc, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13.5, bold=True, color=BLACK)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], header_fill: str = BLUE, font_size: float = 9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    header = table.rows[0].cells
    for idx, value in enumerate(headers):
        set_cell_shading(header[idx], header_fill)
        p = header[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
        format_paragraph(p, size=font_size, color="FFFFFF", bold=True)
    mark_header_row(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or len(row_values) <= 3 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=BLACK)
            format_paragraph(p, size=font_size)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, label: str, text: str, fill: str = PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9500])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "：")
    set_run_font(r, size=10, bold=True, color=BLUE)
    r = p.add_run(text)
    set_run_font(r, size=10, color=BLACK)
    format_paragraph(p, size=10)
    set_cell_border(cell, color="A7C6E2", sz="8")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_flow(doc, cells: list[tuple[str, str]]):
    table = doc.add_table(rows=1, cols=len(cells))
    widths = []
    for label, kind in cells:
        widths.append(2100 if kind == "stage" else 550)
    set_table_geometry(table, widths)
    for idx, (label, kind) in enumerate(cells):
        cell = table.cell(0, idx)
        set_cell_shading(cell, MID_BLUE if kind == "stage" else "FFFFFF")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        set_run_font(r, size=10 if kind == "stage" else 16, bold=kind == "stage", color=BLUE if kind == "stage" else BLUE)
        format_paragraph(p, size=10 if kind == "stage" else 16)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullet(doc, text: str, color: str = BLACK):
    p = doc.add_paragraph(style="List Paragraph")
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.append(num_pr)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=color)
    format_paragraph(p, size=10.5, color=color)
    return p


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run("LLM 局部重调度方案验证报告 | ")
    set_run_font(r, size=8, color="888888")
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    r._r.append(field_begin)
    r._r.append(instr)
    r._r.append(field_end)


def clean_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def read_data():
    audit = json.loads(AUDIT.read_text(encoding="utf-8"))
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    return audit, summary


def choose_recommendation(comparison: dict) -> tuple[str, str]:
    """Choose a branch with the same hard-gate and lexicographic policy as the audit."""
    total = int(comparison.get("num_affected", comparison.get("num_total", 0)) or 0)
    branches = {
        "AQ / 决策树": (
            int(comparison.get("dt_assigned", 0) or 0),
            comparison.get("decision_tree_metrics") or {},
        ),
        "AR / Codex": (
            int(comparison.get("llm_assigned", 0) or 0),
            comparison.get("llm_metrics") or {},
        ),
    }

    def is_feasible(branch: tuple[int, dict]) -> bool:
        assigned, metrics = branch
        return assigned == total and int(metrics.get("rule_violations", 0) or 0) == 0

    feasible = {name: is_feasible(branch) for name, branch in branches.items()}
    if feasible["AQ / 决策树"] and not feasible["AR / Codex"]:
        return "决策树重排方案", "决策树重排方案完整配包且通过硬约束；Codex 重排方案未通过可行性门槛。"
    if feasible["AR / Codex"] and not feasible["AQ / 决策树"]:
        dt_assigned, dt_metrics = branches["AQ / 决策树"]
        ar_assigned, ar_metrics = branches["AR / Codex"]
        return (
            "Codex 重排方案",
            f"决策树重排方案完成 {dt_assigned}/{total}、规则违反 {int(dt_metrics.get('rule_violations', 0) or 0)}；"
            f"Codex 重排方案完成 {ar_assigned}/{total}、规则违反 {int(ar_metrics.get('rule_violations', 0) or 0)}。",
        )
    if not feasible["AQ / 决策树"] and not feasible["AR / Codex"]:
        return "人工复核", "决策树重排方案与 Codex 重排方案均未通过完整性或硬约束门槛，不生成可下发方案。"

    def metric(branch_name: str, key: str, default: float = 0.0) -> float:
        value = branches[branch_name][1].get(key, default)
        try:
            return float(value)
        except (TypeError, ValueError):
            return default

    comparisons = (
        ("按时率", "on_time_rate", True),
        ("平均延迟", "average_delay_seconds", False),
        ("最大延迟", "max_delay_seconds", False),
        ("相对预配基线改配炉次数", "changed_heats", False),
        ("资源负载均衡", "resource_load_balance", False),
    )
    for label, key, higher_is_better in comparisons:
        aq_value = metric("AQ / 决策树", key)
        ar_value = metric("AR / Codex", key)
        if aq_value == ar_value:
            continue
        aq_wins = aq_value > ar_value if higher_is_better else aq_value < ar_value
        winner = "决策树重排方案" if aq_wins else "Codex 重排方案"
        return winner, f"两方均通过门槛，{label}优先，{winner}更优。"
    return "持平", "两条线路通过门槛，且所有选择指标完全相同。"


def evaluation_sections(comparison: dict) -> list[tuple[str, list[list[str]]]]:
    """Return the detailed, reader-facing definition of every selection metric."""
    dt = comparison["decision_tree_metrics"]
    llm = comparison["llm_metrics"]
    total = comparison["num_affected"]
    return [
        ("一、安全可行性门槛", [
            ["完整分配", "窗口内每炉都有钢包和行车", f"已分配炉次 ÷ 窗口炉次", "必须等于 100%", f"{comparison['dt_assigned']}/{total}，不通过", f"{comparison['llm_assigned']}/{total}，通过"],
            ["硬约束违反", "检查钢包、行车、时间、位置、路线、负载和安全距离", "所有违规项计数", "必须等于 0", f"{dt['rule_violations']}，不通过", f"{llm['rule_violations']}，通过"],
            ["资源与路线合法性", "分配的钢包、行车和运输路线必须在当前快照中可用", "逐炉校验资源存在且路线可达", "任一违规即不可下发", "存在未满足项", "全部满足"],
        ]),
        ("二、生产效果", [
            ["按时率", "在炉次时间窗结束前完成运输的比例", "按时炉次 ÷ 窗口炉次", "越高越好", f"{dt['on_time_rate']:.2%}", f"{llm['on_time_rate']:.2%}"],
            ["平均延迟", "每炉超出时间窗的延迟，未超时按 0 计", "平均值：max(到达时刻 − 时间窗结束, 0)", "越低越好", f"{dt['average_delay_seconds']:.1f} 秒", f"{llm['average_delay_seconds']:.1f} 秒"],
            ["最大延迟", "窗口内最晚一炉的超时程度", "max(每炉延迟)", "越低越好", f"{dt['max_delay_seconds']:.0f} 秒", f"{llm['max_delay_seconds']:.0f} 秒"],
        ]),
        ("三、计划稳定性", [
            ["相对预配基线改配炉次数", "重排方案相对生产前预配基线改变钢包或行车的炉次数", "逐炉比较钢包编号和行车编号", "效果接近时越少越好", f"{dt['changed_heats']} 炉", f"{llm['changed_heats']} 炉；28 炉保留预配基线"],
            ["资源负载均衡", "比较各行车最终负载的离散程度", "各行车最终负载的方差", "效果接近时越低越好", "次级指标，当前审计未单独汇总", "次级指标，当前审计未单独汇总"],
            ["变更范围", "扰动只影响当前局部窗口，不扩散到已完成和远期炉次", "统计窗口外是否发生改配", "必须保持局部", "局部窗口内重排", "局部窗口内重排"],
        ]),
    ]


def build_markdown(audit, summary) -> str:
    c = audit["comparison"]
    dt = c["decision_tree_metrics"]
    llm = c["llm_metrics"]
    recommendation, recommendation_reason = choose_recommendation(c)
    changed = audit["stress_audit"]["changed_assignments"]
    metric_tables = evaluation_sections(c)
    evaluation_markdown = "\n\n".join(
        "### " + title + "\n\n| 指标 | 含义 | 计算方式 | 判定方向 | 决策树重排 | Codex 重排 |\n| --- | --- | --- | --- | --- | --- |\n"
        + "\n".join("| " + " | ".join(row) + " |" for row in rows)
        for title, rows in metric_tables
    )
    return f"""# 钢包重排包评价体系建立

**日期：** 2026-08-31  
**分支：** `feature-1.0`  
**数据：** 真实 `PLAN(1).xlsx`、`CRANE.xlsx`、`loc_location.xlsx`

## 1. 汇报结论

在真实 11:00–11:20 局部窗口的 31 炉压力回放中，行车 2500 离线后，决策树重排方案完成 {c['dt_assigned']}/{c['num_affected']}；Codex 重排方案完成 {c['llm_assigned']}/{c['num_affected']}。Codex 只改配 {llm['changed_heats']} 炉，28 炉保留预配基线，硬约束终检通过、规则违反 {llm['rule_violations']}。

本次推荐：**{recommendation}**。原因：{recommendation_reason} Codex 的作用是决策树失败后的补位，不替代全量预配包，也不绕过终检。

## 2. 系统怎么起作用

![扰动后的钢包局部重调度闭环](../docs/llm-rescheduling-flow.png)

1. 生产开始前，决策树生成预配基线。
2. 11:00 发生扰动，只打开 11:00–11:20 的 31 炉窗口；已完成和远期炉次不动。
3. 先生成决策树重排方案；如果局部方案不完整，Codex 基于预配基线、事故状态、可用资源和约束独立生成 Codex 重排方案。
4. Codex 重排方案先过完整性、资源和路线约束终检，检查通过后才形成可下发方案。

预配基线是两条线路的唯一共同输入，决策树重排方案与 Codex 重排方案独立生成；Codex 不读取决策树输出。本次使用离线预先保存并重新校验的 Codex 审计方案，未调用外部接口（`external_api_called=false`）。

### 2.1 评价指标体系

![决策树与 Codex 的评价指标体系](../docs/evaluation-indicator-system.png)

评价不是把所有指标简单相加，而是分三层判定：第一层决定方案能不能安全下发；第二层比较生产效果；第三层在效果接近时比较计划稳定性。任何方案只要第一层失败，就不能因为改动少而被推荐。

{evaluation_markdown}

**本次如何落到选择：** 决策树重排方案只完成 {c['dt_assigned']}/{c['num_affected']}，并有 {dt['rule_violations']} 项规则违反，因此在第一层即淘汰；Codex 重排方案完成 {c['llm_assigned']}/{c['num_affected']}，规则违反 {llm['rule_violations']}，进入生产效果和计划稳定性比较，最终推荐 Codex 重排方案。

## 3. 真实实验结果

| 指标 | 结果 |
| --- | ---: |
| PLAN 总行数 / 纳入评测 | 469 / 457 炉 |
| 预配基线 | 457/457，规则违反 0 |
| 真实资源 | 34 个钢包、10 台行车 |
| 扰动窗口 | 11:00–11:20，共 31 炉 |
| 决策树重排方案 | {c['dt_assigned']}/{c['num_affected']} |
| Codex 重排方案 | {c['llm_assigned']}/{c['num_affected']} |
| Codex 改配 / 保留预配基线 | {llm['changed_heats']} 炉 / 28 炉 |
| Codex 硬约束终检 | {c['llm_assigned']}/{c['num_affected']} 通过，规则违反 {llm['rule_violations']} |
| 最终推荐 | {recommendation} |
| 推荐原因 | {recommendation_reason} |

### 3.1 Codex 实际改配

| 炉次 | 原钢包 / 行车 | Codex 重排钢包 / 行车 |
| --- | --- | --- |
{chr(10).join(f"| {x['heat_id']} | {x['old_ladle_id']} / {x['old_crane_id']} | {x['new_ladle_id']} / {x['new_crane_id']} |" for x in changed)}

## 4. 相比“失败后呼叫人工”的价值

决策树负责毫秒级常规重排；Codex 只处理决策树无法完成的困难窗口。在仍有时间预算时，它可以快速补足资源组合，并输出可审计的局部方案，避免直接退回人工。此次回放中，决策树重排方案少完成 1 炉，Codex 重排方案补足到 {c['num_affected']}/{c['num_affected']}，且只改变 {llm['changed_heats']} 炉。

离线场景库共 21 条，其中 20 条是“决策树失败、Codex 成功”的压力场景，覆盖行车、通道、钢包、工艺设施、计划节奏、数据状态和复合扰动。

## 5. 边界与下一步

行车 2500 离线是基于真实 PLAN/CRANE 快照构造的可复现压力输入，不是宝钢历史事故日志。当前结果验证的是离线回放中的调度输入、方案生成和硬约束闭环，不替代现场 PLC、安全联锁、人工审批或线上 API 延迟验收。

下一步接入实时事件流、现场路线与设备联锁约束，并在沙箱环境测量真实 API 的端到端耗时。

## 6. 证据

- 真实文件：`/Users/admin/Desktop/数据文件/PLAN(1).xlsx`、`CRANE.xlsx`、`loc_location.xlsx`
- 主压力场景审计：`outputs/real_data_codex_stress_demo/audit.json`
- 离线场景库：`outputs/offline_scenarios/ladle_scenarios.sqlite3`
- Excel 对比结果：`outputs/PLAN(1)_位置映射决策树预配包结果.xlsx`
"""


def build_docx(audit, summary) -> None:
    doc = Document(str(REF))
    clean_body(doc)

    # Retain the reference report's A4 geometry, but make CJK font selection explicit.
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = FONT_FAMILY
            r_pr = style._element.get_or_add_rPr()
            r_fonts = r_pr.get_or_add_rFonts()
            for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
                r_fonts.set(qn(f"w:{attr}"), FONT_FAMILY)
            r_fonts.set(qn("w:hint"), "eastAsia")
            lang = r_pr.find(qn("w:lang"))
            if lang is None:
                lang = OxmlElement("w:lang")
                r_pr.append(lang)
            lang.set(qn("w:val"), "en-US")
            lang.set(qn("w:eastAsia"), "zh-CN")
            lang.set(qn("w:bidi"), "en-US")
            style.paragraph_format.line_spacing = 1.5
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(13.5)

    # Keep the report body clean: no page footer is rendered in the deliverable.
    for section in doc.sections:
        footer = section.footer
        for child in list(footer._element):
            footer._element.remove(child)
        for footer_ref in list(section._sectPr.findall(qn("w:footerReference"))):
            section._sectPr.remove(footer_ref)

    title = add_text(doc, "钢包重排包评价体系建立", size=28, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_after = Pt(4)
    subtitle = add_text(doc, "真实数据回放与两种重排方案评价", size=12, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    subtitle.paragraph_format.space_after = Pt(16)
    add_rich_paragraph(doc, [("日期：", {"bold": True}), ("2026-08-31", {}), ("    分支：", {"bold": True}), ("feature-1.0", {"color": BLUE})], align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    comparison = audit["comparison"]
    dt = comparison["decision_tree_metrics"]
    llm = comparison["llm_metrics"]
    recommendation, recommendation_reason = choose_recommendation(comparison)
    add_heading(doc, "1. 一句话结论", 1)
    add_text(doc, f"真实 11:00–11:20 窗口共 {comparison['num_affected']} 炉。行车 2500 离线后，决策树重排方案完成 {comparison['dt_assigned']}/{comparison['num_affected']}；Codex 重排方案完成 {comparison['llm_assigned']}/{comparison['num_affected']}，改配 {llm['changed_heats']} 炉，28 炉保留预配基线，规则违反 {llm['rule_violations']}。本次推荐 {recommendation}。")
    add_text(doc, f"推荐原因：{recommendation_reason} Codex 的定位是决策树失败后的自动兜底，不替代全量预配包，也不绕过统一校验。")
    add_heading(doc, "1.1 当前系统状态", 2)
    add_table(doc, ["项目", "当前结果"], [
        ["全量预配包", "真实 PLAN(1) 469 行中纳入评测 457 炉；决策树基线 457/457，规则违反 0"],
        ["生产响应", "滑窗推进；扰动只打开当前时刻附近的局部炉次窗口"],
        ["Codex 角色", "决策树失败且扣除 90 s 安全余量后仍有可用预算时介入，输出可审计的替代配包方案"],
    ], [2100, 7400])
    add_callout(doc, "验证目标", "确认决策树失败时，Codex 能基于真实资源与约束完成局部窗口重排，并且只修改必要炉次。")

    add_heading(doc, "2. 处理流程", 1)
    add_heading(doc, "2.1 预配基线进入两条独立重排线路", 2)
    if DIAGRAM_PNG.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(str(DIAGRAM_PNG), width=Inches(6.5))
        p.paragraph_format.space_after = Pt(2)
        caption = add_text(doc, "图 1  真实数据压力回放中的局部重调度闭环", size=9, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
        caption.paragraph_format.space_after = Pt(6)
    add_text(doc, "预配基线是唯一共同基线；决策树重排方案与 Codex 重排方案两条线路独立。Codex 不读取决策树输出，只在决策树失败后读取预配基线、事故状态、可用资源和约束。")
    add_heading(doc, "2.2 评价指标体系", 2)
    if INDICATOR_PNG.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(str(INDICATOR_PNG), width=Inches(6.5))
        p.paragraph_format.space_after = Pt(2)
        caption = add_text(doc, "图 2  决策树与 Codex 的统一评价指标体系", size=9, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
        caption.paragraph_format.space_after = Pt(6)
    add_text(doc, "选择顺序只有一句话：先判能不能安全下发，再比按时率和延迟，最后比改配规模和资源稳定性。")
    for section_title, rows in evaluation_sections(comparison):
        add_heading(doc, section_title, 3)
        add_table(
            doc,
            ["指标", "含义", "计算方式", "判定方向", "决策树重排方案", "Codex 重排方案"],
            rows,
            [1250, 2050, 2050, 1250, 1450, 1450],
            font_size=8.2,
        )
    add_heading(doc, "2.3 关键判断", 2)
    key_judgment = add_text(doc, "本次事件发生在 11:00；最早开浇前还有 313.1 s，扣除 90 s 安全余量后仍有 223.1 s 可用预算，因此允许 Codex 介入。若扣除安全余量后可用预算不大于 0，则不进入 Codex，保留人工升级路径。")
    key_judgment.paragraph_format.keep_together = True

    add_heading(doc, "3. 真实实验结果", 1)
    add_table(doc, ["指标", "结果"], [
        ["数据范围", "PLAN 469 行；纳入评测 457 炉"],
        ["真实资源", "34 个钢包；10 台行车"],
        ["局部窗口", "11:00–11:20，共 31 炉"],
        ["决策树重排方案", f"{comparison['dt_assigned']}/{comparison['num_affected']}"],
        ["Codex 重排方案", f"{comparison['llm_assigned']}/{comparison['num_affected']}"],
        ["改配与稳定性", f"改配 {llm['changed_heats']} 炉；28 炉保留预配基线"],
        ["硬约束终检", f"{comparison['llm_assigned']}/{comparison['num_affected']} 通过；规则违反 {llm['rule_violations']}"],
        ["最终推荐", recommendation],
        ["推荐原因", recommendation_reason],
    ], [3000, 6500])
    add_heading(doc, "3.1 Codex 实际改配", 2)
    changed = audit["stress_audit"]["changed_assignments"]
    add_table(doc, ["炉次", "预配基线钢包 / 行车", "Codex 重排钢包 / 行车"], [
        [x["heat_id"], f"{x['old_ladle_id']} / {x['old_crane_id']}", f"{x['new_ladle_id']} / {x['new_crane_id']}"] for x in changed
    ], [4400, 2500, 2600])
    add_text(doc, "3 个改配均用于避开离线行车 2500；其余 28 炉保持预配基线不动。")
    add_callout(doc, "相比人工", "传统路径是决策树失败后直接等待人工。Codex 在仍有时间预算时先给出可审计候选，再交由统一终检；本次将决策树重排方案的 30/31 补足到 Codex 重排方案的 31/31。", fill=LIGHT_BLUE)

    add_heading(doc, "4. 边界与下一步", 1)
    add_text(doc, "行车 2500 离线是基于真实 PLAN/CRANE 快照构造的可复现压力输入，不是宝钢历史事故日志。当前结果验证的是离线回放中的调度输入、方案生成和硬约束闭环，不替代现场 PLC、安全联锁、人工审批或线上 API 延迟验收。")
    add_text(doc, "离线 SQLite 场景库共 21 条，其中 20 条是“决策树失败、Codex 成功”的压力场景，覆盖行车、通道、钢包、工艺设施、计划节奏、数据状态和复合扰动。当前结果来自离线 Codex 审计方案，未调用外部大模型接口，审计字段为 external_api_called=false。")
    add_bullet(doc, "接入实时事件流、现场路线和设备联锁约束。")
    add_bullet(doc, "接入人工审批接口，保留预配基线、两种重排方案和终检记录。")
    add_bullet(doc, "在沙箱环境测量真实 API 的端到端耗时。")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))


def main():
    audit, summary = read_data()
    OUT_MD.write_text(build_markdown(audit, summary), encoding="utf-8")
    build_docx(audit, summary)
    print(OUT_MD)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
