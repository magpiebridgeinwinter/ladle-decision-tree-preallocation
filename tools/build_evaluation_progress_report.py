from __future__ import annotations

import json
import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/Users/admin/Desktop/数据文件/2026-9-1 评价体系初步建立.docx")
AUDIT = ROOT / "outputs/real_data_codex_stress_demo/audit.json"
SUMMARY = ROOT / "outputs/offline_scenarios/summary.json"
OUTPUT = ROOT / "outputs/2026-9-8 钢包重排评价指标调研进展.docx"

_BASE_SPEC = importlib.util.spec_from_file_location("build_project_report", ROOT / "tools/build_project_report.py")
if _BASE_SPEC is None or _BASE_SPEC.loader is None:
    raise RuntimeError("cannot load shared document helpers")
base = importlib.util.module_from_spec(_BASE_SPEC)
_BASE_SPEC.loader.exec_module(base)
# LibreOffice resolves the macOS PostScript family name reliably during PDF rendering.
base.FONT_FAMILY = "Arial Unicode MS"


def configure_styles(doc: Document) -> None:
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        style.font.name = base.FONT_FAMILY
        r_pr = style._element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
            r_fonts.set(qn(f"w:{attr}"), base.FONT_FAMILY)
        r_fonts.set(qn("w:hint"), "eastAsia")
        lang = r_pr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            r_pr.append(lang)
        lang.set(qn("w:val"), "en-US")
        lang.set(qn("w:eastAsia"), "zh-CN")
        lang.set(qn("w:bidi"), "en-US")
        style.paragraph_format.line_spacing = 1.5
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(13.5)


def remove_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        for child in list(footer._element):
            footer._element.remove(child)
        for footer_ref in list(section._sectPr.findall(qn("w:footerReference"))):
            section._sectPr.remove(footer_ref)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.5) -> None:
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    p.paragraph_format.space_after = Pt(1)
    caption_p = base.add_text(doc, caption, size=9, color=base.GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
    caption_p.paragraph_format.space_after = Pt(5)


def metric_rows(comparison: dict) -> list[list[str]]:
    dt = comparison["decision_tree_metrics"]
    llm = comparison["llm_metrics"]
    return [
        ["完整分配率", "能否覆盖窗口内全部炉次", "已分配炉次 / 窗口炉次", "主指标；必须 100%", "30/31", "31/31"],
        ["硬约束违反", "资源、时间、位置、路线与安全规则", "违规项计数", "主指标；必须 0", str(dt["rule_violations"]), str(llm["rule_violations"])],
        ["资源/路线合法性", "钢包、行车存在且路线可达", "逐炉校验资源与路线", "主指标；任一失败淘汰", "有失败项", "全部通过"],
        ["按时率", "窗口结束前完成运输的比例", "按时炉次 / 窗口炉次", "主指标；越高越好", f"{dt['on_time_rate']:.2%}", f"{llm['on_time_rate']:.2%}"],
        ["平均延迟", "超出时间窗的平均时长", "平均 max(实际到达 - 窗口结束, 0)", "主指标；越低越好", f"{dt['average_delay_seconds']:.1f} 秒", f"{llm['average_delay_seconds']:.1f} 秒"],
        ["最大延迟", "最晚一炉的超时程度", "max(单炉延迟)", "主指标；越低越好", f"{dt['max_delay_seconds']:.0f} 秒", f"{llm['max_delay_seconds']:.0f} 秒"],
        ["相对 AP 改配炉次数", "相对预配基线改变的炉次数", "逐炉比较钢包/行车", "补充；效果接近时越少越好", str(dt["changed_heats"]), f"{llm['changed_heats']}（28 炉保留）"],
        ["窗口外改配炉次数", "是否影响已完成和远期炉次", "统计扰动窗口外的变化", "补充；必须 0", "0", "0"],
        ["资源负载均衡", "方案对各行车负载的影响", "负载方差或最大-最小差", "补充；效果接近时比较", "待补充汇总", "待补充汇总"],
    ]


def build_report() -> None:
    audit = json.loads(AUDIT.read_text(encoding="utf-8"))
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    comparison = audit["comparison"]
    dt = comparison["decision_tree_metrics"]
    llm = comparison["llm_metrics"]

    doc = Document(str(TEMPLATE))
    base.clean_body(doc)
    configure_styles(doc)
    remove_footer(doc)

    title = base.add_text(doc, "钢包重排评价指标调研进展", size=28, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_after = Pt(4)
    subtitle = base.add_text(doc, "本周结论：确定哪些指标适合作为方案评价标准", size=12, color=base.BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    subtitle.paragraph_format.space_after = Pt(14)
    base.add_rich_paragraph(
        doc,
        [("日期：", {"bold": True}), ("2026-09-08", {}), ("    数据：", {"bold": True}), ("PLAN(1) / CRANE / loc_location", {"color": base.BLUE})],
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    base.add_heading(doc, "1. 本周结论", 1)
    base.add_text(doc, "本周完成了评价指标的筛选与真实数据回放验证。最终不采用简单加权总分，而采用“硬门槛 + 分层排序”：先判断方案能不能安全下发，再比较生产效果，最后用计划稳定性做平局裁决。这样可以避免“改配少”掩盖“方案不可执行”。")
    base.add_callout(doc, "已确定", "主评价指标为完整分配率、硬约束违反、资源/路线合法性、按时率、平均延迟和最大延迟；改配规模、窗口外影响和负载均衡作为稳定性补充指标。", fill=base.LIGHT_BLUE)
    base.add_table(doc, ["层级", "要回答的问题", "本周结论"], [
        ["第一层：可行性", "方案能不能安全下发？", "完整分配必须 100%；硬约束必须 0；资源与路线必须合法。"],
        ["第二层：生产效果", "可行方案谁更少延误？", "按时率优先，其次平均延迟、最大延迟。"],
        ["第三层：计划稳定性", "效果接近时谁改动更小？", "比较 AP 基线改配、窗口外影响和资源负载。"],
    ], [1800, 3000, 4700], font_size=9)

    base.add_heading(doc, "2. 指标筛选方法", 1)
    base.add_heading(doc, "2.1 为什么不使用综合总分", 2)
    base.add_text(doc, "钢包重排首先是一个安全可行性问题，不是单纯的成本优化问题。若某方案少改了几炉，但有一炉没有合法钢包、行车或路线，就不能因为加权分数较高而被下发。因此，评价顺序固定为：可行性门槛 → 生产效果 → 稳定性。任一方案未通过第一层，直接淘汰；两方都未通过时转人工复核。")
    add_figure(doc, ROOT / "docs/evaluation-indicator-system.png", "图 1  本周确定的三层评价逻辑")
    base.add_heading(doc, "2.2 推荐指标清单", 2)
    base.add_table(doc, ["指标", "指标含义", "计算方式", "定位", "决策树", "Codex"], metric_rows(comparison), [1250, 1900, 1900, 1550, 1450, 1450], font_size=7.8)

    base.add_heading(doc, "3. 真实数据验证", 1)
    base.add_text(doc, "验证采用真实 PLAN(1)、CRANE 和 loc_location 快照。以 11:00 行车 2500 受控离线为扰动输入，只打开 11:00–11:20 的局部窗口，共 31 炉；已完成炉次和远期炉次不参与重排。该场景是基于真实数据构造的可复现压力回放，不是宝钢历史事故日志。")
    base.add_table(doc, ["验证项", "结果", "对指标筛选的意义"], [
        ["数据规模", "PLAN 469 行；纳入评测 457 炉", "具备完整的基线和资源快照，可计算覆盖率与改配规模。"],
        ["局部扰动", "窗口 31 炉；行车 2500 离线", "能验证局部性和窗口外不扩散。"],
        ["决策树重排", "30/31；规则违反 1 项；按时率 96.77%", "证明仅看改配数量不够，必须保留硬门槛。"],
        ["Codex 重排", "31/31；规则违反 0；按时率 100%", "证明完整性、硬约束和生产效果可以区分两条线路。"],
        ["计划稳定性", "Codex 改配 3 炉；28 炉保留 AP 基线", "改配规模适合作为可行方案之间的次级比较项。"],
    ], [2200, 2600, 4700], font_size=8.6)
    base.add_callout(doc, "本次证据", "决策树方案在第一层就因 30/31 且有 1 项规则违反被淘汰；Codex 方案进入第二、三层比较并通过终检。该结果支持评价顺序，但不代表单一场景足以估计长期成功率。")

    base.add_heading(doc, "3.1 LLM 兜底价值如何评价", 2)
    base.add_text(doc, "Codex 的价值不是在所有炉次上替代决策树，而是在决策树无法完成局部重排时提供第二条可审计线路。因此，LLM 相关指标应跨场景统计，而不是从一次回放推导结论。")
    base.add_table(doc, ["LLM 价值指标", "建议定义", "当前状态"], [
        ["决策树失败后 Codex 成功率", "决策树失败场景中，Codex 通过全部终检的比例", "适合纳入；场景库已有 20 条压力场景"],
        ["人工升级减少率", "相对决策树直接升级人工，Codex 成功接管的比例", "适合纳入；需补齐人工升级基线"],
        ["在线决策耗时 / P95", "从事件到可审计方案的真实端到端耗时", "适合纳入；当前离线 decision_seconds 不作结论"],
        ["输出可审计性", "是否记录输入快照、约束检查、方案和终检结果", "作为交付要求，不与生产效果合并打分"],
    ], [2800, 4000, 2700], font_size=8.4)

    base.add_heading(doc, "4. 暂不作为主指标的项目", 1)
    base.add_text(doc, "以下字段不是没有价值，而是当前数据或实验条件不足以支撑主评价结论。先标记为暂缓，可以避免报告给出超过证据边界的判断。")
    base.add_table(doc, ["暂缓指标", "暂缓原因", "后续补充"], [
        ["单次离线决策耗时", "离线执行的 0 秒不等于真实 API 延迟", "接入真实接口，记录平均值和 P95/P99"],
        ["人工复核率", "当前场景库没有现场人工处理时长和升级记录", "建立人工基线与升级日志"],
        ["资源负载方差", "当前双线路审计尚未统一汇总各行车负载", "按窗口和全天分别统计"],
        ["包龄利用率、等级匹配率", "更适合预配包质量；扰动重排缺少现场状态校准", "补充钢包状态、等级和包龄真值"],
        ["简单综合评分", "会让低改配掩盖不可行方案，解释性也较弱", "保留分层排序，不作为主判定"],
    ], [2600, 4500, 2400], font_size=8.5)
    base.add_heading(doc, "4.1 场景覆盖情况", 2)
    groups = "；".join(f"{x['disturbance_group']} {x['count']} 条" for x in summary["disturbance_groups"])
    base.add_text(doc, f"离线场景库共 {summary['scenario_count']} 条：1 条决策树成功对照、20 条决策树失败后 Codex 成功的压力场景。扰动覆盖 {groups}。这些场景用于验证指标和兜底链路的可计算性，不能替代现场事故统计。")
    base.add_callout(doc, "评价口径", "推荐顺序固定为：完整性 → 硬约束 → 按时率 → 平均延迟 → 最大延迟 → 改配规模 → 负载均衡。若两套方案均未过硬门槛，则不比较总分，直接人工复核。", fill=base.PALE_BLUE)

    base.add_heading(doc, "5. 下一步调研计划", 1)
    base.add_table(doc, ["工作项", "目的", "输出"], [
        ["补齐现场状态字段", "统一钢包状态、位置、路线、设备联锁和包龄口径", "可复现实验快照"],
        ["建立多场景对照集", "让决策树成功、失败和人工升级都有基线", "按场景统计成功率与 P95"],
        ["接入真实 LLM 接口计时", "测量 90 秒安全预算内是否可完成", "平均耗时、P95、超时率"],
        ["形成验收规则", "将硬门槛和分层排序固化到审计与 Excel 输出", "自动评价报告"],
    ], [2600, 4400, 2500], font_size=8.6)
    base.add_text(doc, "本周产出是评价体系的初版口径：先保证方案可用，再比较效果，最后控制改动。后续随着现场事件和人工处理数据补齐，可把 LLM 兜底成功率、人工减少率和真实时延纳入正式验收。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    sys.path.insert(0, str(ROOT))
    build_report()
