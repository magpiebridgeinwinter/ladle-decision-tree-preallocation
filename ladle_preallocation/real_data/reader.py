"""Read production workbooks in read-only mode and align rows to package contracts."""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from ladle_preallocation.data_modeling.alignment import align_crane_row, align_master_row
from ladle_preallocation.data_modeling.assumptions import DEFAULT_X_LOWER_LIMIT, DEFAULT_X_UPPER_LIMIT
from ladle_preallocation.data_modeling.features import parse_time

CRANE_SHEET = "CRANE_STATUS_HIS - 副本"
PLAN_SHEET = "PLAN_TAPPING - 副本"
LOCATION_SHEET = "Sheet1"
PLAN_MACHINE_HEADERS = [f"Column{i}" for i in range(1, 53)]
PLAN_LOCATION_HEADERS = {
    "mapped_span",
    "mapped_pos_x",
    "mapped_loc_type",
    "mapped_loc_status",
    "mapped_station_name",
    "location_mapping_status",
}
LOCATION_HEADERS = ("LOC_LOCATION", "SPAN_NAME", "POS_X", "LOC_TYPE", "LOC_STATUS", "GWDMC")


def parse_production_time(value: Any) -> tuple[float | None, bool]:
    """Accept the specified microsecond format and PLAN's observed seconds-only variant."""
    timestamp, invalid = parse_time(value)
    if not invalid:
        return timestamp, False
    try:
        return datetime.strptime(str(value), "%Y-%m-%d-%H.%M.%S").replace(tzinfo=timezone.utc).timestamp(), False
    except (TypeError, ValueError):
        return None, True


def _values(ws: Any, start_row: int, max_rows: int | None) -> Iterable[tuple[Any, ...]]:
    emitted = 0
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if not any(value not in (None, "") for value in row):
            continue
        yield row
        emitted += 1
        if max_rows is not None and emitted >= max_rows:
            break


def read_crane(path: str | Path, max_rows: int | None = None) -> list[dict[str, Any]]:
    """Read rows from CRANE.xlsx; blank package IDs remain ``None`` by contract."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if CRANE_SHEET not in workbook.sheetnames:
            raise ValueError(f"未找到行车工作表: {CRANE_SHEET}")
        sheet = workbook[CRANE_SHEET]
        if sheet.max_column != 20:
            raise ValueError(f"行车数据列序不一致：期望20列，实际{sheet.max_column}列")
        return [align_crane_row(row) for row in _values(sheet, 2, max_rows)]
    finally:
        workbook.close()


def read_latest_online_cranes_in_window(path: str | Path, start_at: float, end_at: float) -> list[dict[str, Any]]:
    """Stream CRANE history and retain only the latest online snapshot per crane.

    This avoids materialising the million-row history when pre-allocation needs
    only the operating state covering the filtered PLAN time window.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if CRANE_SHEET not in workbook.sheetnames:
            raise ValueError(f"未找到行车工作表: {CRANE_SHEET}")
        sheet = workbook[CRANE_SHEET]
        if sheet.max_column != 20:
            raise ValueError(f"行车数据列序不一致：期望20列，实际{sheet.max_column}列")
        latest: dict[str, tuple[float, dict[str, Any]]] = {}
        for row in _values(sheet, 2, None):
            record = align_crane_row(row)
            timestamp, invalid = parse_production_time(record.get("updated_at"))
            if invalid or timestamp is None or timestamp < start_at or timestamp > end_at:
                continue
            if str(record.get("online_status")) not in {"1", "1.0"}:
                continue
            crane_id = str(record["crane_id"])
            if crane_id not in latest or timestamp > latest[crane_id][0]:
                latest[crane_id] = (timestamp, record)
        return [item[1] for item in latest.values()]
    finally:
        workbook.close()


def read_plan(path: str | Path, max_rows: int | None = None) -> list[dict[str, Any]]:
    """Read a PLAN workbook, including optional in-sheet location mapping columns."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if PLAN_SHEET not in workbook.sheetnames:
            raise ValueError(f"未找到炉次工作表: {PLAN_SHEET}")
        sheet = workbook[PLAN_SHEET]
        machine_headers = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
        if machine_headers == PLAN_MACHINE_HEADERS:
            return [align_master_row(row, machine_headers) for row in _values(sheet, 3, max_rows)]
        if not machine_headers or any(header not in PLAN_MACHINE_HEADERS and header not in PLAN_LOCATION_HEADERS for header in machine_headers):
            raise ValueError("列序不一致：PLAN 第2行必须为 ColumnN 标签，可附加位置映射列")
        # A filtered export retains the original ColumnN labels. Expand it before
        # alignment so business meaning is never inferred from its new position.
        column_indexes = [int(header.removeprefix("Column")) - 1 if header in PLAN_MACHINE_HEADERS else None for header in machine_headers]
        records = []
        for row in _values(sheet, 3, max_rows):
            full_row = [None] * 52
            mapped_values = {}
            for header, index, value in zip(machine_headers, column_indexes, row):
                if index is None:
                    mapped_values[header] = value
                else:
                    full_row[index] = value
            record = align_master_row(full_row, PLAN_MACHINE_HEADERS)
            record.update(mapped_values)
            record["_source_pre_filtered"] = True
            records.append(record)
        return records
    finally:
        workbook.close()


def read_location(path: str | Path, max_rows: int | None = None) -> dict[str, dict[str, Any]]:
    """Read the plant location dictionary keyed by its stable location code."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if LOCATION_SHEET not in workbook.sheetnames:
            raise ValueError(f"未找到位置字典工作表: {LOCATION_SHEET}")
        sheet = workbook[LOCATION_SHEET]
        header = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        if header[: len(LOCATION_HEADERS)] != LOCATION_HEADERS:
            raise ValueError(
                f"位置字典列序不一致：期望 {LOCATION_HEADERS}，实际 {header[:len(LOCATION_HEADERS)]}"
            )
        locations: dict[str, dict[str, Any]] = {}
        emitted = 0
        for row in _values(sheet, 2, max_rows):
            values = list(row[: len(LOCATION_HEADERS)])
            if not any(value not in (None, "") for value in values):
                continue
            code = str(values[0] or "").strip()
            if not code:
                raise ValueError("位置字典存在空的 LOC_LOCATION")
            if code in locations:
                raise ValueError(f"位置字典存在重复 LOC_LOCATION: {code}")
            raw_x = values[2]
            try:
                pos_x = None if raw_x in (None, "") else float(raw_x)
            except (TypeError, ValueError):
                pos_x = None
            position_valid = (
                pos_x is not None
                and DEFAULT_X_LOWER_LIMIT <= pos_x <= DEFAULT_X_UPPER_LIMIT
            )
            locations[code] = {
                "location_code": code,
                "span_name": values[1],
                "pos_x": pos_x,
                "loc_type": values[3],
                "loc_status": values[4],
                "description": values[5],
                "position_valid": position_valid,
                "position_invalid_reason": (
                    None if position_valid else (
                        "missing_or_non_numeric" if pos_x is None else "outside_coordinate_range"
                    )
                ),
            }
            emitted += 1
            if max_rows is not None and emitted >= max_rows:
                break
        return locations
    finally:
        workbook.close()


def location_mapping_statistics(
    plan_records: list[dict[str, Any]], location_map: dict[str, dict[str, Any]]
) -> dict[str, Any]:
    """Summarize how PLAN position codes resolve against the location dictionary."""
    counts = Counter()
    for row in plan_records:
        code = str(row.get("ladle_position") or "").strip()
        if not code:
            counts["missing_plan_code"] += 1
            continue
        record = location_map.get(code)
        if record is None:
            counts["missing_location_code"] += 1
        elif not record.get("position_valid"):
            counts["missing_coordinate"] += 1
        else:
            counts["mapped"] += 1
    counts["total_plan_rows"] = len(plan_records)
    counts["unique_plan_codes"] = len({str(row.get("ladle_position") or "").strip() for row in plan_records if row.get("ladle_position") not in (None, "")})
    counts["dictionary_rows"] = len(location_map)
    return dict(counts)


def dataset_statistics(records: list[dict[str, Any]], critical_fields: tuple[str, ...], time_field: str) -> dict[str, Any]:
    """Return compact, serializable volume, missingness, and time-range statistics."""
    total = len(records)
    missing = {field: sum(row.get(field) in (None, "") for row in records) / total if total else 0.0 for field in critical_fields}
    timestamps = [value for value, invalid in (parse_production_time(row.get(time_field)) for row in records) if not invalid and value is not None]
    return {"total_rows": total, "valid_rows": sum(any(value not in (None, "") for value in row.values()) for row in records), "missing_rate": missing, "time_range_unix": [min(timestamps), max(timestamps)] if timestamps else None}


def grade_distribution(plan_records: list[dict[str, Any]]) -> dict[str, int]:
    values = [str(row.get("required_grade")).strip().upper() for row in plan_records if row.get("required_grade") not in (None, "")]
    categories = Counter("numeric" if value in {"0", "1", "4", "5"} else "alpha_or_other" for value in values)
    return {**dict(categories), **{f"grade_{value}": values.count(value) for value in ("0", "1", "4", "5")}}
